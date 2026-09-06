"""
应用制作托管（公开路由，token 门控）

「应用制作」任务：agent 在 users/{uid}/apps/{workflow_key}/ 目录开发多文件 HTML 应用，
前端 iframe 以 `/api/v1/ai/html-app/{token}/index.html` 为 src 渲染。

为什么走 URL token 而不是 Authorization 头：
- iframe 的 src 请求无法自定义请求头；页面内相对引用的 css/js/子页面与 fetch('./data.json') 也都不带头。
- 因此 token 内嵌在 URL 路径里，所有相对引用自动解析到同一 token 前缀下（token 即凭据）。

token 语义（app/utils/security.py::create_html_app_token）：
- HS256 无状态签名，TTL 7d（分享态 1d），claims = {tokenType: "htmlAppToken", userId, workflowKey[, share][, visitorId]}
  （visitorId 仅「仅登录用户」分享模式嵌入访客身份；免登录分享访客保持匿名）
- tokenType 与登录态双向隔离：dependency.py 硬校验 tokenType == "accessToken"，本 token 打任何
  鉴权接口都被拒；反向本路由只认 htmlAppToken，登录 token 也不能当托管 token 用。
- 权限域 = 单个任务目录的「读任意文件 + 写 json + 用户上传文件落盘与文档文本提取」
  + 受 SSRF 门控的公网数据代理 + 行数据通道（agent_app_row 大表，$acl 行级保护）+ 平台模型 AI 对话通道；
  删任务即删目录与行数据，已签发 token 自然全部 404。
- share=true（看板开启分享后访客经 share-view 端点取得；分享分「仅登录用户 / 免登录」双模式）：
  数据通道读写能力与板主一致（$acl 可按身份限制到表）；仅 serve html 时不注入「编辑文字」脚本
  （不允许编辑看板本身）；/ai 的模型配置与计费归因跟随生效身份（仅登录 → 访客本人，免登录 → 板主）。

红线（对照 CLAUDE.md）：
- 公开端点绝不使用 4001/4002/4003/4010 鉴权契约码（那是 axios 拦截器的登出/刷新信号）；
  GET（iframe 消费）失败返回 PlainTextResponse 404，POST save（页面 JS 消费）失败用 Fail 4000 系。
- async 端点内所有磁盘 I/O 走 asyncio.to_thread。
- token 会出现在服务端访问日志（URL 内嵌的固有代价，内网管理系统可接受）；
  Referer 泄漏由 iframe referrerpolicy + 页面 meta referrer no-referrer 兜住。
"""

import asyncio
import base64
import io
import ipaddress
import json
import os
import re
import socket
import time
from collections import deque
from pathlib import Path
from typing import Any, Optional
from urllib.parse import urlparse

import httpx
from fastapi import APIRouter, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, PlainTextResponse, StreamingResponse
from loguru import logger
from pydantic import BaseModel
from tortoise.transactions import in_transaction

from app.core.ctx import CTX_BILLING_BIZ_ENTRY, CTX_USER_ID
from app.models.standard.agent import AgentAppRow
from app.schemas.base import Fail, Success
from app.services.agent_runtime import app_events
from app.utils.security import decode_html_app_token

router = APIRouter(prefix="/html-app", tags=["应用制作托管（公开）"])

# workspace 根（与 agent_public.py / qa.py 同款约定）
_PROJECT_ROOT = Path(__file__).parent.parent.parent.parent.parent
_WORKSPACE = _PROJECT_ROOT / ".agent_workspace"
_USERS_ROOT = _WORKSPACE / "users"

# 单文件写回上限（json 数据层定位，防把应用中转站当网盘）
_SAVE_MAX_BYTES = 2 * 1024 * 1024

# 外部数据代理参数（防 SSRF / 防把服务器当开放中转）
_PROXY_TIMEOUT = 60.0
_PROXY_REQ_MAX_BYTES = 1 * 1024 * 1024
_PROXY_RESP_MAX_BYTES = 5 * 1024 * 1024
_PROXY_METHODS = {"GET", "POST", "PUT", "DELETE", "PATCH"}
# 业务头白名单（小写匹配）；host / cookie 等身份类头绝不转发
_PROXY_HEADER_ALLOW = {
    "content-type", "accept", "accept-language", "authorization",
    "x-api-key", "apikey", "api-key", "x-auth-token", "user-agent",
}
# 这些 content-type 按文本直出；其余按 base64 返回（页面侧按 encoding 字段区分）
_PROXY_TEXT_PREFIXES = ("application/json", "text/", "application/xml", "application/javascript")

# 页面 AI 对话通道参数（防把入口当无限转发的刷量器）
_AI_MAX_MESSAGES = 50
_AI_MAX_CHARS_PER_MSG = 30_000
_AI_TIMEOUT = 600.0
# /ai 附图（多模态直传 chat 模型，由属主 chat 块的 vision 能力门控）：
# 限额服务于 DashScope 6MB 请求体上限——raw 合计 3MB → base64 ≈ 4MB，给对话文本留 2MB
_AI_MAX_IMAGES_TOTAL = 3
_AI_IMAGE_MAX_BYTES = 2 * 1024 * 1024
_AI_IMAGES_TOTAL_MAX_BYTES = 3 * 1024 * 1024
# /ai 限流（进程内滑动窗口）：run.py 单 worker 单事件循环，窗口检查与记账之间无 await，
# 普通 dict + deque 即天然原子；将来 workers>1 再换 Redis 分桶（依赖已有）
_AI_RATE_WINDOW_S = 60.0
_AI_RATE_LIMIT_IDENTIFIED = 10  # 登录身份：每应用每用户 10 次/分
_AI_RATE_LIMIT_ANONYMOUS = 6  # 匿名访客：每应用每客户端 IP 6 次/分
_ai_rate_windows: dict[str, deque[float]] = {}


def _ai_rate_limited(bucket: str, limit: int) -> bool:
    """滑动窗口限流：命中返回 True（不记账）；未命中记一次并返回 False。"""
    now = time.monotonic()
    q = _ai_rate_windows.get(bucket)
    if q is None:
        q = deque()
        _ai_rate_windows[bucket] = q
    while q and now - q[0] >= _AI_RATE_WINDOW_S:
        q.popleft()
    if len(q) >= limit:
        return True
    q.append(now)
    # 懒清理：桶数膨胀时按批淘汰已过期桶（只在此处顺带做，摊还 O(1)）
    if len(_ai_rate_windows) > 4096:
        stale = [k for k, v in list(_ai_rate_windows.items()) if not v or now - v[-1] >= _AI_RATE_WINDOW_S][:1024]
        for k in stale:
            _ai_rate_windows.pop(k, None)
    return False


# 文件上传通道参数（用户本地文件 → 任务目录 uploads/；定位是应用的业务文件，不是网盘）
# 标准 multipart 上传；解析期 starlette 先把文件 spool 到临时文件（>1MB 自动落盘，不吃内存），
# 端点内再流式拷贝进任务目录并边写边计数，超限即中止清半成品——限额同时是 nginx 侧配置的依据
_UPLOAD_MAX_BYTES = 10 * 1024 * 1024 * 1024
_UPLOAD_CHUNK = 4 * 1024 * 1024
# 扩展名白名单：文档类 + 图片 + 常见媒体/压缩包；可执行/代码类一律拒（应用自身代码由 agent 写入，不走此通道）
_UPLOAD_EXTS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".rtf",
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log",
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".ico",
    ".mp3", ".wav", ".mp4", ".webm", ".zip",
}
# 文本提取：返回正文上限（超长截断 + textTruncated，页面自行分段喂 /ai）；源文件过大直接拒绝解析
_EXTRACT_MAX_CHARS = 200_000
_EXTRACT_MAX_SOURCE_BYTES = 50 * 1024 * 1024
# 可按文本直读的扩展名（utf-8 容错解码）；docx / xlsx / pdf 走专用解析
_PLAIN_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log"}

# 自带 MIME 映射（不依赖 mimetypes：Windows 注册表可能污染/缺失映射，中文文本类型必须带 charset）
_EXT_MIME = {
    ".html": "text/html", ".htm": "text/html",
    ".js": "text/javascript", ".mjs": "text/javascript",
    ".css": "text/css",
    ".json": "application/json",
    ".svg": "image/svg+xml",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".webp": "image/webp", ".avif": "image/avif", ".ico": "image/x-icon",
    ".woff": "font/woff", ".woff2": "font/woff2", ".ttf": "font/ttf", ".otf": "font/otf",
    ".mp4": "video/mp4", ".webm": "video/webm", ".mp3": "audio/mpeg", ".wav": "audio/wav",
    ".txt": "text/plain", ".md": "text/plain", ".csv": "text/csv", ".tsv": "text/tab-separated-values",
    ".log": "text/plain", ".xml": "application/xml",
    ".wasm": "application/wasm", ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".ppt": "application/vnd.ms-powerpoint",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".rtf": "application/rtf", ".zip": "application/zip",
}
_TEXT_MIME = {"text/html", "text/javascript", "text/css", "application/json", "text/plain", "text/csv"}

_NO_STORE = {"Cache-Control": "no-store"}

# ── 系统注入的「编辑文字」能力 ─────────────────────────────────────────────────
# 响应 html 页面时自动注入 html_app_edit.js（不落盘、不进存档）：入口是画布外框顶栏的
# 「编辑文字」按钮（workflow/index.vue），宿主与页内脚本走 postMessage；用户就地改文本，改动以
# {from: 原文, to: 改后} 替换对经 /save 存进 data/text-edits.json，页面加载时自动套用。
# 只改文本不碰样式/结构；agent 迭代页面时负责合并这份清单（见 HTML_BOARD_RULES）。
# 分享访客（share token）以只读模式注入（window.__hbteReadOnly）：保存过的文字改动照常自动套用
# （否则访客看到的永远是没有编辑过的原版），但编辑开关被忽略、写回清单也被 save 通道拒绝。
_EDIT_SCRIPT_PATH = Path(__file__).parent / "html_app_edit.js"
_edit_script_cache: Optional[str] = None

# ── 系统注入的「互动桥」window.AppBridge ───────────────────────────────────────
# 与编辑脚本同机制（不落盘、不进存档），但**无条件注入所有 html 页面**（板主 / 访客皆然）：
# 封装行数据读写（loadRows/putRows/delRows）、SSE 变更订阅（onChange）、反向呼叫 agent
# （sendToChat）。页面可只用它而不手写 fetch。脚本内容不含 "</script>" 字样，裸拼安全；
# 文件缺失则静默降级（只注入编辑脚本），存量应用零影响。
_BRIDGE_SCRIPT_PATH = Path(__file__).parent / "html_app_bridge.js"
_bridge_script_cache: Optional[str] = None


def _load_edit_script() -> str:
    """读取注入脚本（首次读后进程内缓存）；文件缺失返回空串 = 能力静默关闭，不影响页面托管。"""
    global _edit_script_cache
    if _edit_script_cache is None:
        try:
            _edit_script_cache = _EDIT_SCRIPT_PATH.read_text(encoding="utf-8")
        except OSError:
            _edit_script_cache = ""
    return _edit_script_cache


def _load_bridge_script() -> str:
    """读取互动桥脚本（首次读后进程内缓存）；文件缺失返回空串 = 能力静默关闭。"""
    global _bridge_script_cache
    if _bridge_script_cache is None:
        try:
            _bridge_script_cache = _BRIDGE_SCRIPT_PATH.read_text(encoding="utf-8")
        except OSError:
            _bridge_script_cache = ""
    return _bridge_script_cache


def _maybe_inject_edit_script(target: Path, file_path: str, read_only: bool = False) -> Optional[HTMLResponse]:
    """serve 的是 html 页面 → 在最后一个 </body> 前注入「互动桥 + 编辑」两个脚本并返回 HTMLResponse；
    其余文件返回 None 走原路。

    注入顺序：AppBridge 无条件在前（所有页面可用），编辑脚本在后（维持 read_only 语义）。
    read_only=True（分享访客）：编辑脚本先置 window.__hbteReadOnly——页面加载时照常套用已保存的
    文字替换对，但编辑开关消息被忽略、改动也不会写回；AppBridge 不受 read_only 影响（访客照样
    能读写行数据 / 订阅变更，sendToChat 在分享页无宿主监听会静默降级 no-channel）。"""
    if not file_path.lower().endswith((".html", ".htm")):
        return None
    bridge = _load_bridge_script()
    edit = _load_edit_script()
    if not bridge and not edit:
        return None
    try:
        text = target.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None
    # 注入的是裸 JS，必须包 <script> 标签才会执行；脚本内容本身不含 "</script>" 字样，直接拼安全
    tags: list[str] = []
    if bridge:
        tags.append("<script>\n" + bridge + "\n</script>")
    if edit:
        flag = "window.__hbteReadOnly = true;\n" if read_only else ""
        tags.append("<script>\n" + flag + edit + "\n</script>")
    block = "\n".join(tags)
    pos = text.lower().rfind("</body>")
    injected = text[:pos] + "\n" + block + "\n" + text[pos:] if pos != -1 else text + "\n" + block
    return HTMLResponse(injected, headers=_NO_STORE)


def _app_base(user_id: int, workflow_key: str) -> Path:
    return _USERS_ROOT / str(user_id) / "apps" / workflow_key


def _resolve_inside(base: Path, rel: str) -> Optional[Path]:
    """把相对路径解析到 base 内；任何越权/非法形态返回 None。

    三重拒绝 + 双 relative_to 保险：`..` 段、反斜杠（%5C）、绝对路径/盘符先拒；
    再 resolve 后分别对 base 与 users 根做 relative_to，symlink 逃逸也兜住。
    """
    if not rel or "\\" in rel or rel.startswith("/") or ":" in rel.split("/")[0]:
        return None
    parts = rel.split("/")
    if ".." in parts or "." in parts:
        return None
    # 点目录 / 点文件一律拒：应用目录内的 .versions/（发布存档）等系统数据不得被页面读写
    # （合法应用资源与 uploads 文件名都不会以点开头——上传名清洗已拒点开头）
    if any(seg.startswith(".") for seg in parts):
        return None
    try:
        target = (base / rel).resolve()
        target.relative_to(base.resolve())
        target.relative_to(_USERS_ROOT.resolve())
    except (ValueError, OSError):
        return None
    return target


def _media_type(path: Path) -> str:
    mt = _EXT_MIME.get(path.suffix.lower(), "application/octet-stream")
    if mt in _TEXT_MIME:
        mt += "; charset=utf-8"
    return mt


def _write_atomic(target: Path, text: str) -> None:
    """同目录 tmp + os.replace 原子替换（Windows NTFS 上同样原子，避免 iframe 读到半截 JSON）。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, target)


# ── 行数据通道（agent_app_row 大表）─────────────────────────────────────────────
# 应用的用户数据统一层：一行 = 应用里的一条业务记录，(workflow_key, tbl, row_key) 三元组定位。
# tbl 是应用自定义的逻辑集合名（如 members / records）；`$` 开头为平台保留——$acl 表存
# 行级保护清单（被保护表名 = row_key，data = {writers, readers}），强制规则见 _acl_allows。
_ACL_TBL = "$acl"
_TBL_RE = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_.-]{0,63}$")
_ROW_KEY_MAX = 191
_ROWS_BATCH_MAX = 200
_ROW_DATA_MAX_BYTES = 256 * 1024
_ROWS_QUERY_LIMIT_MAX = 1000
_ROWS_QUERY_LIMIT_DEFAULT = 200


def _valid_tbl(tbl: str) -> bool:
    """应用侧逻辑集合名：字母/数字/下划线开头，可含 . _ -，最长 64 字符；$ 开头平台保留（正则天然排除）。"""
    return isinstance(tbl, str) and bool(_TBL_RE.match(tbl))


def _valid_row_key(key: str) -> bool:
    """行键：1~191 字符任意串（与表列宽一致，应用自定语义，多用户按人数据建议 {visitorId}: 前缀）。"""
    return isinstance(key, str) and 1 <= len(key) <= _ROW_KEY_MAX


def _acl_allows(entry: Any, actor_uid: int | None, owner_uid: int, mode: str) -> bool:
    """$acl 强制判定（纯函数，便于单测）。

    板主恒过；无规则 / 规则畸形一律放行（缺省放行：应用未配置保护时行为不变）；
    匿名访客（actor_uid=None）永不命中 uid 列表；mode = "read" / "write" 对应 readers / writers 字段，
    该字段缺省（None）= 不限制该向。
    """
    if actor_uid is not None and actor_uid == owner_uid:
        return True
    if not isinstance(entry, dict):
        return True
    rule = entry.get("writers" if mode == "write" else "readers")
    if rule is None:
        return True
    if not isinstance(rule, list):
        return True
    return actor_uid in rule


def _validate_acl_data(data: Any) -> Optional[str]:
    """$acl 行数据形状校验：{writers?: [uid...], readers?: [uid...]}。返回错误描述；None = 通过。

    uid 元素接受数字或数字字符串（whoami 返回字符串 uid，应用可能原样塞入），
    通过后 data 内的名单已就地归一为纯数字（强制判定 _acl_allows 按数字比较）。"""
    if not isinstance(data, dict):
        return "必须是 JSON 对象"
    extra = set(data.keys()) - {"writers", "readers"}
    if extra:
        return f"含未识别字段：{', '.join(sorted(extra))}"
    for field in ("writers", "readers"):
        lst = data.get(field)
        if lst is None:
            continue
        if not isinstance(lst, list):
            return f"{field} 必须是用户ID数组"
        norm: list[int] = []
        for x in lst:
            if isinstance(x, bool):
                return f"{field} 必须是用户ID数组"
            if isinstance(x, int):
                norm.append(x)
            elif isinstance(x, str) and x.isdigit():
                norm.append(int(x))
            else:
                return f"{field} 必须是用户ID数组"
        data[field] = norm
    return None


async def _load_tbl_acl(workflow_key: str) -> dict[str, dict]:
    """一次查询载入该应用的 $acl 保护清单：{被保护表名: {"writers": [...], "readers": [...]}}。

    畸形行数据（非 dict）就地跳过 = 该表无保护（与 _acl_allows 的缺省放行一致）。"""
    rows = await AgentAppRow.filter(workflow_key=workflow_key, tbl=_ACL_TBL).all()
    acl: dict[str, dict] = {}
    for r in rows:
        if isinstance(r.data, dict):
            acl[r.row_key] = r.data
    return acl


def _actor_from_claims(claims: dict) -> tuple[int, Optional[int], bool]:
    """从 token claims 解析 (owner_uid, actor_uid, is_owner)。

    板主 token（无 share 声明）→ actor = 板主；「仅登录」分享 token（带 visitorId）→ actor = 访客；
    免登录分享 token → actor = None（匿名，永不命中 $acl uid 名单）。"""
    owner_uid = int(claims["userId"])
    if not claims.get("share"):
        return owner_uid, owner_uid, True
    visitor_uid = claims.get("visitorId")
    return owner_uid, visitor_uid if isinstance(visitor_uid, int) else None, False


@router.get("/{token}/whoami", summary="应用内当前身份（token 门控）")
async def whoami(token: str):
    """返回 token claims 派生的身份信息，供应用实现多用户 / 角色逻辑（绝不让用户自报身份）。

    ownerId = 板主 uid；visitorId = 「仅登录用户」分享模式下的访客 uid（免登录 / 板主 token 为 null）；
    isOwner 由「无 share 声明」推得（板主 token）。nickName / userName 尽力而为取当前行动者
    （有访客取访客，否则板主）的昵称与账号名，查不到为 null。

    uid 一律以**字符串**返回：行数据通道的行键是字符串，uid 给字符串后应用可直接用作行键、
    直接 === 比较，无需 String()/Number() 转换（历史教训：返回数字时 LLM 生成的应用普遍
    忘转换，导致成员注册被校验拒绝、身份与行键比较永不命中）。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    owner_uid, actor_uid, is_owner = _actor_from_claims(claims)
    nick_name: Optional[str] = None
    user_name: Optional[str] = None
    try:
        from app.models.system.admin import User

        u = await User.get_or_none(id=actor_uid if actor_uid is not None else owner_uid)
        if u:
            nick_name = u.nick_name or u.user_name
            user_name = u.user_name
    except Exception as e:  # noqa: BLE001 —— 尽力而为字段，失败仍返回身份主体
        logger.warning(f"[html_app] whoami 用户查询失败 uid={actor_uid}: {e!r}")

    return Success(
        data={
            "ownerId": str(owner_uid),
            "visitorId": None if (is_owner or actor_uid is None) else str(actor_uid),
            "isOwner": is_owner,
            "share": bool(claims.get("share")),
            "nickName": nick_name,
            "userName": user_name,
        }
    )


@router.get("/{token}/rows", summary="读取应用行数据（token 门控，$acl readers 强制）")
async def list_rows(token: str, tbl: str, prefix: Optional[str] = None, limit: int = _ROWS_QUERY_LIMIT_DEFAULT, offset: int = 0):
    """按行键排序分页读取某逻辑集合：{rows: [{key, data, updatedBy, updatedAt}], more}。

    $acl 表所有人可读（应用渲染权限 UI 用）；其余表按 $acl 的 readers 名单强制。
    updatedBy（最后写入者 uid）与行键同口径返回字符串，便于应用直接 === 比较。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    tbl = (tbl or "").strip()
    if tbl != _ACL_TBL and not _valid_tbl(tbl):
        return Fail(code="4000", msg="tbl 名不合法")
    if not 1 <= limit <= _ROWS_QUERY_LIMIT_MAX:
        return Fail(code="4000", msg=f"limit 需在 1~{_ROWS_QUERY_LIMIT_MAX}")
    if offset < 0:
        return Fail(code="4000", msg="offset 不能为负")

    owner_uid, actor_uid, _is_owner = _actor_from_claims(claims)
    key = str(claims["workflowKey"])

    if tbl != _ACL_TBL:
        acl = await _load_tbl_acl(key)
        if not _acl_allows(acl.get(tbl), actor_uid, owner_uid, "read"):
            logger.warning(f"[html_app] rows 读被拒 key={key} tbl={tbl} actor={actor_uid}")
            return Fail(code="4000", msg="没有该数据表的读取权限")

    qs = AgentAppRow.filter(workflow_key=key, tbl=tbl)
    if prefix:
        qs = qs.filter(row_key__startswith=prefix)
    rows = await qs.order_by("row_key").offset(offset).limit(limit + 1)
    more = len(rows) > limit
    return Success(
        data={
            "rows": [
                {
                    "key": r.row_key,
                    "data": r.data,
                    "updatedBy": str(r.updated_by) if r.updated_by is not None else None,
                    "updatedAt": r.update_time.isoformat() if r.update_time else None,
                }
                for r in rows[:limit]
            ],
            "more": more,
        }
    )


class HtmlAppRowItem(BaseModel):
    # 行键兼容纯数字（归一成字符串）：whoami 的 uid 虽是字符串，应用仍可能把未转换的
    # 数字 uid 直通作行键，这里宽容归一，避免成员注册类写入整批被校验拒绝
    key: str | int
    data: Any


class HtmlAppRowsPut(BaseModel):
    tbl: str
    rows: list[HtmlAppRowItem]


@router.post("/{token}/rows/put", summary="写入应用行数据（token 门控，$acl writers 强制）")
async def put_rows(token: str, body: HtmlAppRowsPut):
    """upsert 语义：同 (tbl, key) 覆盖，否则插入。≤200 行/批，单行 data dumps 后 ≤256KB。

    $acl 表仅板主可写（行键 = 被保护的表名，数据形状 {writers?, readers?}，uid 数字或数字字符串均可、
    归一成数字入库）；其余表按 writers 名单强制。行键为纯数字时自动转字符串。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    tbl = (body.tbl or "").strip()
    is_acl = tbl == _ACL_TBL
    if not is_acl and not _valid_tbl(tbl):
        return Fail(code="4000", msg="tbl 名不合法")
    rows = body.rows or []
    if not rows or len(rows) > _ROWS_BATCH_MAX:
        return Fail(code="4000", msg=f"rows 条数需在 1~{_ROWS_BATCH_MAX}")
    for item in rows:
        if isinstance(item.key, int) and not isinstance(item.key, bool):
            item.key = str(item.key)
        if not _valid_row_key(item.key):
            return Fail(code="4000", msg=f"行键不合法：{(item.key or '')[:32]}")
        if len(json.dumps(item.data, ensure_ascii=False).encode("utf-8")) > _ROW_DATA_MAX_BYTES:
            return Fail(code="4000", msg=f"行数据过大：{item.key[:32]}")

    owner_uid, actor_uid, _is_owner = _actor_from_claims(claims)
    key = str(claims["workflowKey"])

    if is_acl:
        if actor_uid != owner_uid:
            logger.warning(f"[html_app] $acl 写被拒 key={key} actor={actor_uid}")
            return Fail(code="4000", msg="仅板主可修改权限配置")
        for item in rows:
            if not _valid_tbl(item.key):
                return Fail(code="4000", msg=f"$acl 行键必须是合法表名：{item.key[:32]}")
            err = _validate_acl_data(item.data)
            if err:
                return Fail(code="4000", msg=f"$acl 行数据不合法：{err}")
    else:
        acl = await _load_tbl_acl(key)
        if not _acl_allows(acl.get(tbl), actor_uid, owner_uid, "write"):
            logger.warning(f"[html_app] rows 写被拒 key={key} tbl={tbl} actor={actor_uid}")
            return Fail(code="4000", msg="没有该数据表的写入权限")

    async with in_transaction():
        for item in rows:
            await AgentAppRow.update_or_create(
                workflow_key=key,
                tbl=tbl,
                row_key=item.key,
                defaults={"data": item.data, "updated_by": actor_uid},
            )
    logger.info(f"[html_app] rows/put key={key} tbl={tbl} n={len(rows)} actor={actor_uid}")
    # 行数据变更 → SSE 广播（by=page）：订阅该板的所有页面（板主 + 分享访客）onChange 刷新
    app_events.publish(key, [tbl], "page")
    return Success(data={"written": len(rows)})


class HtmlAppRowsDel(BaseModel):
    tbl: str
    # 与写入同口径：行键兼容纯数字（归一成字符串）
    keys: list[str | int]


@router.post("/{token}/rows/del", summary="删除应用行数据（token 门控，$acl writers 强制）")
async def del_rows(token: str, body: HtmlAppRowsDel):
    """按行键批量删除，≤200 键/批。$acl 表仅板主可删；其余表按 writers 名单强制。行键为纯数字时自动转字符串。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    tbl = (body.tbl or "").strip()
    is_acl = tbl == _ACL_TBL
    if not is_acl and not _valid_tbl(tbl):
        return Fail(code="4000", msg="tbl 名不合法")
    keys = body.keys or []
    if not keys or len(keys) > _ROWS_BATCH_MAX:
        return Fail(code="4000", msg=f"keys 条数需在 1~{_ROWS_BATCH_MAX}")
    keys = [str(k) if isinstance(k, int) and not isinstance(k, bool) else k for k in keys]
    if any(not _valid_row_key(k) for k in keys):
        return Fail(code="4000", msg="行键不合法")

    owner_uid, actor_uid, _is_owner = _actor_from_claims(claims)
    key = str(claims["workflowKey"])

    if is_acl:
        if actor_uid != owner_uid:
            logger.warning(f"[html_app] $acl 删被拒 key={key} actor={actor_uid}")
            return Fail(code="4000", msg="仅板主可修改权限配置")
    else:
        acl = await _load_tbl_acl(key)
        if not _acl_allows(acl.get(tbl), actor_uid, owner_uid, "write"):
            logger.warning(f"[html_app] rows 删被拒 key={key} tbl={tbl} actor={actor_uid}")
            return Fail(code="4000", msg="没有该数据表的写入权限")

    deleted = await AgentAppRow.filter(workflow_key=key, tbl=tbl, row_key__in=keys).delete()
    logger.info(f"[html_app] rows/del key={key} tbl={tbl} n={deleted} actor={actor_uid}")
    app_events.publish(key, [tbl], "page")
    return Success(data={"deleted": deleted})


# ── SSE 行数据变更推送（token 门控）────────────────────────────────────────────
# ⚠ 注册顺序红线：必须在下面的 catch-all serve_file（/{token}/{file_path:path}）之前，
# 否则 /events 会被当成静态文件路径吞掉。EventSource 无法自定义请求头 → token 走 URL
# （与页面所有请求同口径）；失败返回 PlainTextResponse 404，绝不碰 4001/4002/4003/4010 契约码。
# 载荷最小化（不带行数据）：订阅者含匿名分享访客，行数据可能受 $acl readers 保护，
# 页面收到事件后自行走 loadRows 正规读通道刷新（丢中间事件无害，幂等）。
_SSE_HEARTBEAT_S = 20.0


@router.get("/{token}/events", summary="订阅应用行数据变更（SSE，token 门控）")
async def row_events(token: str):
    claims = decode_html_app_token(token)
    if claims is None:
        return PlainTextResponse("链接无效或已过期", status_code=404)
    wf_key = str(claims["workflowKey"])

    q = app_events.subscribe(wf_key)
    if q is None:
        # 订阅配额打满（连接泄漏 / 恶意刷）：503 让 EventSource 走原生重连退避
        return PlainTextResponse("订阅数已满，请稍后重试", status_code=503)

    async def gen():
        try:
            # 首帧：立即发一个 hello 事件，页面据此确认通道就绪（也冲掉代理层的首包缓冲）
            yield "event: hello\ndata: {}\n\n"
            while True:
                try:
                    ev = await asyncio.wait_for(q.get(), timeout=_SSE_HEARTBEAT_S)
                    yield f"event: rows\ndata: {json.dumps(ev, ensure_ascii=False)}\n\n"
                except asyncio.TimeoutError:
                    # 心跳注释帧：防代理/浏览器判定连接空闲而断开
                    yield ": ping\n\n"
        except asyncio.CancelledError:
            # 客户端断开（EventSource close / 页面卸载）：正常收尾，finally 清理订阅
            raise
        finally:
            app_events.unsubscribe(wf_key, q)

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={**_NO_STORE, "X-Accel-Buffering": "no", "Connection": "keep-alive"},
    )


@router.get("/{token}/{file_path:path}", summary="托管应用制作静态文件（token 门控）")
async def serve_file(token: str, file_path: str):
    claims = decode_html_app_token(token)
    if claims is None:
        return PlainTextResponse("链接无效或已过期", status_code=404)

    # 目录形态兜底：/ 或 pages/ 结尾 → index.html
    if not file_path:
        file_path = "index.html"
    elif file_path.endswith("/"):
        file_path += "index.html"

    base = _app_base(int(claims["userId"]), str(claims["workflowKey"]))
    target = _resolve_inside(base, file_path)
    if target is None:
        return PlainTextResponse("链接无效或已过期", status_code=404)

    exists = await asyncio.to_thread(target.is_file)
    if not exists:
        # 跨机惰性物化：本机没有该板文件而 DB 有已发布版本（另一台实例发布的）→ 捞回磁盘再试。
        # 失败不阻塞，退回原 404 语义
        try:
            from app.api.v1.ai.agent_workflow import ensure_app_files

            if await ensure_app_files(int(claims["userId"]), str(claims["workflowKey"])):
                exists = await asyncio.to_thread(target.is_file)
        except Exception:  # noqa: BLE001 —— 物化是增强路径，任何异常保持 404 原行为
            pass
    if not exists:
        return PlainTextResponse("文件不存在", status_code=404)

    # 分享访客（share token）以只读模式注入编辑脚本：已保存的文字改动自动套用（访客看到板主编辑后的样子），
    # 但脚本忽略编辑开关、save 通道也拒写文字清单——「不能编辑看板本身」不变
    injected = await asyncio.to_thread(_maybe_inject_edit_script, target, file_path, bool(claims.get("share")))
    if injected is not None:
        return injected
    return FileResponse(target, media_type=_media_type(target), headers=_NO_STORE)


class HtmlAppSave(BaseModel):
    path: str
    data: Any

    class Config:
        populate_by_name = True


@router.post("/{token}/save", summary="页面内写回 json 数据（token 门控）")
async def save_json(token: str, body: HtmlAppSave):
    """应用页面的数据持久化通道：仅允许任务目录内的 .json 文件，服务端 dumps 保证合法 JSON。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    rel = (body.path or "").strip()
    if not rel.lower().endswith(".json"):
        return Fail(code="4000", msg="仅支持保存 .json 数据文件")
    # 文字编辑清单是「板主编辑看板」的落点：share token 只读套用、禁止写回（防访客绕过只读注入改板）
    if claims.get("share") and rel.replace("\\", "/").lower() == "data/text-edits.json":
        return Fail(code="4000", msg="分享访问不支持修改看板文字")

    base = _app_base(int(claims["userId"]), str(claims["workflowKey"]))
    target = _resolve_inside(base, rel)
    if target is None:
        return Fail(code="4000", msg="非法路径")

    text = json.dumps(body.data, ensure_ascii=False)
    raw = text.encode("utf-8")
    if len(raw) > _SAVE_MAX_BYTES:
        return Fail(code="4000", msg="数据过大")

    await asyncio.to_thread(_write_atomic, target, text)
    # 双写 DB 活动集：用户数据的跨机一致真相源。本机落盘已成功，DB 写失败只记日志
    # （最坏退回单机一致，与入库机制上线前行为相同），不向页面报错
    try:
        from app.api.v1.ai.agent_workflow import db_upsert_live_file

        await db_upsert_live_file(str(claims["workflowKey"]), rel, raw)
    except Exception as e:  # noqa: BLE001 —— 数据已在本机落盘，DB 同步失败不阻塞页面
        logger.warning(f"[html_app] save 双写 DB 失败 {claims['workflowKey']}/{rel}: {e!r}")
    return Success(msg="已保存")


def _ip_blocked(ip: ipaddress.IPv4Address | ipaddress.IPv6Address) -> bool:
    """内网 / 环回 / 链路本地 / 多播 / 保留 / 未指定地址一律拒绝（SSRF 门控）。"""
    return (
        ip.is_private
        or ip.is_loopback
        or ip.is_link_local
        or ip.is_multicast
        or ip.is_reserved
        or ip.is_unspecified
    )


def _proxy_target_blocked(url: str) -> Optional[str]:
    """代理目标安全校验：仅放行 http/https 且解析结果全部为公网地址。返回拒绝原因；None = 放行。

    解析出的**每一个**地址都要检查（多 A 记录里混一个内网地址也拒），防止内网探测 / 打本机服务。
    """
    try:
        parsed = urlparse(url)
    except ValueError:
        return "非法 URL"
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return "仅支持 http/https 地址"
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    try:
        infos = socket.getaddrinfo(parsed.hostname, port, proto=socket.IPPROTO_TCP)
    except (socket.gaierror, OSError):
        return "域名解析失败"
    if not infos:
        return "域名解析失败"
    for info in infos:
        try:
            if _ip_blocked(ipaddress.ip_address(info[4][0])):
                return "目标地址不在允许范围"
        except ValueError:
            return "目标地址不在允许范围"
    return None


class HtmlAppProxy(BaseModel):
    url: str
    method: str = "GET"
    headers: Optional[dict[str, str]] = None
    body: Optional[str] = None


@router.post("/{token}/proxy", summary="页面外部数据代理（token 门控，仅公网 http/https）")
async def proxy_request(token: str, body: HtmlAppProxy):
    """应用的外部数据通道：iframe 页面跨域 fetch 会被 CORS 拦，改由服务端代发。

    响应统一信封 {status, contentType, body, truncated}：文本类 content-type body 为 utf-8 字符串，
    其余为 base64（多一个 encoding='base64' 字段）；上游 3xx 不跟随，原样返回由页面处理。
    """
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    method = (body.method or "GET").upper()
    if method not in _PROXY_METHODS:
        return Fail(code="4000", msg="仅支持 GET/POST/PUT/DELETE/PATCH")

    url = (body.url or "").strip()
    blocked = await asyncio.to_thread(_proxy_target_blocked, url)
    if blocked:
        return Fail(code="4000", msg=f"拒绝代理：{blocked}")

    payload = body.body.encode("utf-8") if body.body is not None else None
    if payload is not None and len(payload) > _PROXY_REQ_MAX_BYTES:
        return Fail(code="4000", msg="请求体过大")

    headers = {k: v for k, v in (body.headers or {}).items() if k.lower() in _PROXY_HEADER_ALLOW}
    try:
        async with httpx.AsyncClient(follow_redirects=False, timeout=_PROXY_TIMEOUT) as client:
            resp = await client.request(method, url, headers=headers, content=payload)
    except httpx.HTTPError:
        return Fail(code="4000", msg="外部请求失败（超时或网络不可达）")

    # 多读 1 字节判断截断；超限部分丢弃
    raw = resp.content[: _PROXY_RESP_MAX_BYTES + 1]
    truncated = len(raw) > _PROXY_RESP_MAX_BYTES
    raw = raw[:_PROXY_RESP_MAX_BYTES]

    ctype = resp.headers.get("content-type", "")
    base_type = ctype.split(";")[0].strip().lower()
    data: dict[str, Any] = {
        "status": resp.status_code,
        "contentType": ctype,
        "truncated": truncated,
    }
    if not base_type or base_type.startswith(_PROXY_TEXT_PREFIXES):
        data["body"] = raw.decode("utf-8", "replace")
    else:
        data["body"] = base64.b64encode(raw).decode("ascii")
        data["encoding"] = "base64"
    return Success(data=data)


class HtmlAppAiMessage(BaseModel):
    role: str
    content: str
    images: Optional[list[str]] = None  # 可选附图：图片字节的 base64，仅 user 消息可带


class HtmlAppAi(BaseModel):
    messages: list[HtmlAppAiMessage]


@router.post("/{token}/ai", summary="页面 AI 对话（token 门控，走平台 chat 模型）")
async def ai_chat(token: str, body: HtmlAppAi, request: Request):
    """应用的 AI 能力通道：让应用本身成为 AI 应用（对话 / 生成 / 分析），无需用户提供任何 key。

    生效身份 = 「仅登录用户」分享（token 带 visitorId）→ 访客，否则板主（免登录分享跟随板主）；
    模型跟随生效身份的按角色模型配置（chat_block_key），计费与配额也归因到生效身份
    （经 CTX_USER_ID / CTX_BILLING_BIZ_ENTRY，每请求独立 context，set 后无需 reset）。
    限流：登录身份 10 次/分/应用+用户，匿名 6 次/分/应用+IP。
    入参为 OpenAI 风格 messages（system/user/assistant），user 消息可带 images（base64 数组）
    走多模态直传——由生效 chat 块的 vision 能力门控，纯文本块直接 4000 拒绝
    （绝不塌缩重试，页面据此提示用户）。
    """
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    msgs = body.messages or []
    if not msgs or len(msgs) > _AI_MAX_MESSAGES:
        return Fail(code="4000", msg=f"messages 条数需在 1~{_AI_MAX_MESSAGES}")
    if sum(len(m.content or "") for m in msgs) > _AI_MAX_CHARS_PER_MSG * 2:
        return Fail(code="4000", msg="输入内容过长")

    owner_uid, actor_uid, _is_owner = _actor_from_claims(claims)
    wf_key = str(claims["workflowKey"])

    # 限流（先于配额 / 模型解析，防刷量穿透到下游）
    if actor_uid is not None:
        bucket, rate_limit = f"{wf_key}|u{actor_uid}", _AI_RATE_LIMIT_IDENTIFIED
    else:
        client_ip = request.client.host if request.client else "?"
        bucket, rate_limit = f"{wf_key}|ip{client_ip}", _AI_RATE_LIMIT_ANONYMOUS
    if _ai_rate_limited(bucket, rate_limit):
        logger.warning(f"[html_app] ai 限流命中 key={wf_key} actor={actor_uid} bucket={bucket}")
        return Fail(code="4000", msg="请求过于频繁，请稍后再试")

    # 计费 / 配额归因到生效身份：匿名访客的消耗记板主（免登录分享跟随板主），登录访客记访客本人
    effective_uid = actor_uid if actor_uid is not None else owner_uid
    CTX_USER_ID.set(effective_uid)
    CTX_BILLING_BIZ_ENTRY.set("html-app")

    from app.langchain.billing.quota import check_quota

    quota_status = await check_quota(effective_uid)
    if not quota_status.allowed:
        return Fail(code="4000", msg="积分余额不足，请联系管理员充值")

    try:
        from app.langchain.llm_providers import get_chat_llm_for_block, get_llm
        from app.langchain.role_model_profile import resolve_user_model_profile

        profile = await resolve_user_model_profile(effective_uid)
        llm = get_chat_llm_for_block(profile.chat_block_key) if profile.chat_block_key else get_llm()
    except Exception as e:  # noqa: BLE001 —— 页面 JS 只消费 msg，细节留服务端日志
        logger.warning(f"[html_app] ai_chat 模型配置解析失败 user={effective_uid}: {e!r}")
        return Fail(code="4000", msg="AI 配置加载失败，请稍后重试")

    from langchain_core.messages import AIMessage, HumanMessage, SystemMessage

    from app.api.v1.ai.upload import _sniff_image_mime

    lc_messages: list = []
    image_count = 0
    image_bytes_total = 0
    for m in msgs:
        role = (m.role or "").strip().lower()
        images = m.images or []
        if images and role != "user":
            return Fail(code="4000", msg="图片附件只能放在 user 消息上")
        if role == "system":
            lc_messages.append(SystemMessage(content=m.content))
            continue
        if role == "assistant":
            lc_messages.append(AIMessage(content=m.content))
            continue
        if role != "user":
            return Fail(code="4000", msg=f"不支持的角色：{m.role}")
        if not images:
            lc_messages.append(HumanMessage(content=m.content))
            continue
        # 带图 user 消息：vision 能力门控（绑定生效身份实际生效的 chat 块，无回退）+ 逐张校验
        if not profile.supports_vision:
            return Fail(code="4000", msg="当前聊天模型不支持图片理解，请去掉图片附件")
        image_count += len(images)
        if image_count > _AI_MAX_IMAGES_TOTAL:
            return Fail(code="4000", msg=f"最多附带 {_AI_MAX_IMAGES_TOTAL} 张图片")
        content_blocks: list = [{"type": "text", "text": m.content or ""}]
        for b64 in images:
            b64 = (b64 or "").strip()
            try:
                raw = base64.b64decode(b64, validate=True)
            except ValueError:
                return Fail(code="4000", msg="images 必须是合法的 base64")
            if len(raw) > _AI_IMAGE_MAX_BYTES:
                return Fail(code="4000", msg=f"单张图片过大（上限 {_AI_IMAGE_MAX_BYTES // 1048576}MB）")
            image_bytes_total += len(raw)
            if image_bytes_total > _AI_IMAGES_TOTAL_MAX_BYTES:
                return Fail(code="4000", msg=f"图片总大小超限（上限 {_AI_IMAGES_TOTAL_MAX_BYTES // 1048576}MB）")
            mime = _sniff_image_mime(raw)
            if mime is None:
                return Fail(code="4000", msg="仅支持 png / jpg / webp / gif / bmp 图片")
            content_blocks.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}})
        lc_messages.append(HumanMessage(content=content_blocks))

    try:
        result = await asyncio.wait_for(llm.ainvoke(lc_messages), timeout=_AI_TIMEOUT)
    except asyncio.TimeoutError:
        return Fail(code="4000", msg="AI 响应超时，请稍后重试")
    except Exception as e:  # noqa: BLE001 —— 页面 JS 只消费 msg，细节留服务端日志
        logger.warning(f"[html_app] ai_chat 失败 user={effective_uid}: {e!r}")
        return Fail(code="4000", msg="AI 调用失败，请稍后重试")

    content = result.content
    if isinstance(content, list):  # thinking 系模型返回块列表，只拼文本块
        content = "".join(
            part.get("text", "") if isinstance(part, dict) else str(part)
            for part in content
            if not isinstance(part, dict) or part.get("type", "text") == "text"
        )
    return Success(data={"content": content})


# ── 文件通道：用户上传落盘 + 文档文本提取 ─────────────────────────────────────
# /upload 是环境里唯一的 multipart 端点（标准 FormData，文件大、走流式不膨胀）；
# 其余端点（save / proxy / ai / extract）仍一律 JSON。上传落进任务目录 uploads/
# （静态 GET 原样可读回，删任务随之清除）；/extract 给目录内已有文件
# （agent 写入的附件、先前上传的文件）提取正文喂 /ai。


def _sanitize_upload_name(name: str) -> Optional[str]:
    """上传文件名清洗：只留 basename、去掉路径与危险字符。返回 None = 非法。"""
    name = (name or "").replace("\\", "/").strip()
    name = name.split("/")[-1]
    name = re.sub(r"[\x00-\x1f<>:\"|?*]+", "_", name).strip(" ._")
    if not name or name.startswith("."):
        return None
    if len(name) > 120:  # 超长保扩展名截断
        stem, dot, ext = name.rpartition(".")
        name = stem[: max(1, 120 - len(ext) - 1)] + dot + ext if dot else name[:120]
    return name


def _unique_filename(directory: Path, name: str) -> str:
    """同名追加 _N 后缀（与 upload.py 同款约定）。"""
    stem = Path(name).stem
    suffix = Path(name).suffix
    candidate = name
    counter = 1
    while (directory / candidate).exists():
        candidate = f"{stem}_{counter}{suffix}"
        counter += 1
    return candidate


def _save_upload_stream(base: Path, name: str, src, limit: int) -> tuple[Optional[str], int]:
    """从 spool 临时文件流式拷贝进 {base}/uploads/，边写边计数（10GB 级文件不整读进内存）。

    返回 (最终文件名, 已写字节数)；超限返回 (None, 已写数)，半成品在 finally 清掉。
    """
    up_dir = base / "uploads"
    up_dir.mkdir(parents=True, exist_ok=True)
    final = _unique_filename(up_dir, name)
    tmp = up_dir / (final + ".tmp")
    written = 0
    try:
        with open(tmp, "wb") as f:
            while True:
                chunk = src.read(_UPLOAD_CHUNK)
                if not chunk:
                    break
                written += len(chunk)
                if written > limit:
                    return None, written
                f.write(chunk)
        os.replace(tmp, up_dir / final)
        return final, written
    finally:
        if tmp.exists():  # os.replace 成功后 tmp 已不存在；超限/异常清掉半成品
            tmp.unlink()


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _extract_docx(raw: bytes) -> str:
    """docx = zip 包里的 word/document.xml；按段落聚合 w:t 文本（std lib，零依赖）。"""
    import zipfile
    from xml.etree import ElementTree as ET

    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        xml_bytes = zf.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    lines: list[str] = []
    for p in root.iter(_W_NS + "p"):
        text = "".join(t.text or "" for t in p.iter(_W_NS + "t"))
        if text.strip():
            lines.append(text)
    return "\n".join(lines)


def _extract_xlsx(raw: bytes) -> str:
    """xlsx 逐表逐行转 TSV 文本（openpyxl 只读流式，data_only 取公式计算值）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 20_000:
                    lines.append("（行数过多，已截断）")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    lines.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


def _extract_pdf(raw: bytes) -> str:
    """pypdf 逐页提取（加密 PDF 尝试空口令，失败抛错由上层兜成 error）。"""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if reader.is_encrypted:
        reader.decrypt("")
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        if i >= 500:
            parts.append("（页数过多，已截断）")
            break
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


_EXTRACTORS = {".docx": ("docx", _extract_docx), ".xlsx": ("xlsx", _extract_xlsx), ".pdf": ("pdf", _extract_pdf)}


def _safe_extract(name: str, raw: bytes) -> tuple[str, str, bool]:
    """按扩展名提取正文，返回 (text, kind, truncated)，绝不抛异常。

    kind：text / docx / xlsx / pdf / unsupported（格式不支持）/ error（文件损坏或加密）。
    页面 JS 只消费这三个字段，解析失败 = text 为空 + kind 说明原因，由页面如实提示用户。
    """
    ext = Path(name).suffix.lower()
    try:
        if ext in _PLAIN_EXTS:
            text, kind = raw.decode("utf-8", "replace"), "text"
        elif ext in _EXTRACTORS:
            kind, fn = _EXTRACTORS[ext]
            text = fn(raw)
        else:
            return "", "unsupported", False
    except Exception as e:  # noqa: BLE001 —— 损坏/加密/畸形文件统一兜底
        logger.warning(f"[html_app] 文本提取失败 {name}: {type(e).__name__}: {e}")
        return "", "error", False
    truncated = len(text) > _EXTRACT_MAX_CHARS
    return text[:_EXTRACT_MAX_CHARS], kind, truncated


@router.post("/{token}/upload", summary="页面上传用户本地文件到任务目录（token 门控）")
async def upload_file(
    token: str,
    file: UploadFile = File(...),
    extractText: bool = Form(False),
):
    """应用的文件入口：标准 multipart 上传（FormData，file 字段），单文件上限 10GB。

    落盘到任务目录 uploads/ 下（同名自动加后缀），页面之后可 fetch('uploads/xxx') 读回；
    extractText=true 时顺带提取正文文本返回（供喂 /ai 分析，支持 pdf / docx / xlsx / 纯文本系；
    超过 _EXTRACT_MAX_SOURCE_BYTES 的大文件跳过提取，textKind=too_large）。
    """
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    name = _sanitize_upload_name(file.filename or "")
    if not name:
        return Fail(code="4000", msg="文件名不合法")
    ext = Path(name).suffix.lower()
    if ext not in _UPLOAD_EXTS:
        return Fail(code="4000", msg=f"不支持上传该类型文件（{ext or '无扩展名'}）")

    base = _app_base(int(claims["userId"]), str(claims["workflowKey"]))
    try:
        final, size = await asyncio.to_thread(_save_upload_stream, base, name, file.file, _UPLOAD_MAX_BYTES)
    finally:
        await file.close()
    if final is None:
        return Fail(code="4000", msg=f"文件过大（上限 {_UPLOAD_MAX_BYTES // (1024 ** 3)}GB）")
    if size == 0:
        (base / "uploads" / final).unlink(missing_ok=True)
        return Fail(code="4000", msg="文件内容为空")

    data: dict[str, Any] = {"path": f"uploads/{final}", "size": size}
    if extractText:
        if size <= _EXTRACT_MAX_SOURCE_BYTES:
            raw = await asyncio.to_thread((base / "uploads" / final).read_bytes)
            text, kind, truncated = await asyncio.to_thread(_safe_extract, final, raw)
        else:
            text, kind, truncated = "", "too_large", False
        data.update({"text": text, "textKind": kind, "textChars": len(text), "textTruncated": truncated})
    return Success(data=data)


class HtmlAppExtract(BaseModel):
    path: str


@router.post("/{token}/extract", summary="提取任务目录内文件的正文文本（token 门控）")
async def extract_text(token: str, body: HtmlAppExtract):
    """目录内已有文件的文本提取：agent 写入的附件、先前上传的文件，无需重新上传即可解析喂 /ai。"""
    claims = decode_html_app_token(token)
    if claims is None:
        return Fail(code="4030", msg="链接无效或已过期")

    rel = (body.path or "").strip()
    base = _app_base(int(claims["userId"]), str(claims["workflowKey"]))
    target = _resolve_inside(base, rel)
    if target is None:
        return Fail(code="4000", msg="非法路径")
    if not await asyncio.to_thread(target.is_file):
        # 跨机惰性物化兜底（与静态托管同机制）：本机缺文件而 DB 有已发布版本 → 捞回磁盘再试
        try:
            from app.api.v1.ai.agent_workflow import ensure_app_files

            if await ensure_app_files(int(claims["userId"]), str(claims["workflowKey"])):
                pass
        except Exception:  # noqa: BLE001 —— 物化是增强路径，失败保持原报错语义
            pass
    if not await asyncio.to_thread(target.is_file):
        return Fail(code="4000", msg="文件不存在")

    size = (await asyncio.to_thread(target.stat)).st_size
    if size > _EXTRACT_MAX_SOURCE_BYTES:
        return Fail(code="4000", msg="文件过大，无法解析")

    raw = await asyncio.to_thread(target.read_bytes)
    text, kind, truncated = await asyncio.to_thread(_safe_extract, target.name, raw)
    return Success(data={"path": rel, "size": size, "text": text, "textKind": kind, "textChars": len(text), "textTruncated": truncated})
